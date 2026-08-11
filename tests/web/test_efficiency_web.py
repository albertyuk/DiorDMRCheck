"""Web layer for the KOL efficiency report: upload → report view → download.

Verifies the privacy contract too: no upload lands on disk or in the DB, and
expired tokens 404.
"""
from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient

from app import config
from app import main as main_mod
from app.efficiency import routes as eff_routes
from tests.fixtures import build_eff_bytes


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(config, "APP_PASSWORD", "")   # open mode — no login
    eff_routes.EFF_REPORTS.clear()
    with TestClient(main_mod.app) as c:
        yield c
    eff_routes.EFF_REPORTS.clear()


def _upload(client, **form):
    return client.post(
        "/efficiency",
        files={"report": ("wave1.xlsx", build_eff_bytes(),
                          "application/vnd.openxmlformats-officedocument"
                          ".spreadsheetml.sheet")},
        data=form, follow_redirects=True)


def test_full_flow(client, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    r = _upload(client)
    assert r.status_code == 200
    body = r.text
    assert "Download .pptx" in body
    assert "MID PAID" in body and "V10" in body            # findings shown WITH output
    assert "CPM WINNER" in body
    # nothing was written to disk — client campaign data stays in memory
    assert not (tmp_path / "uploads").exists()

    token = next(iter(eff_routes.EFF_REPORTS))
    d = client.get(f"/efficiency/{token}/deck.pptx")
    assert d.status_code == 200
    assert d.content[:2] == b"PK"
    assert "wave1_efficiency.pptx" in d.headers["content-disposition"]


def test_expired_token_404(client):
    _upload(client)
    token = next(iter(eff_routes.EFF_REPORTS))
    eff_routes.EFF_REPORTS[token]["created"] -= eff_routes.EFF_REPORTS.ttl_seconds + 1
    assert client.get(f"/efficiency/{token}").status_code == 404
    assert client.get(f"/efficiency/{token}/deck.pptx").status_code == 404


def test_blocked_report_has_no_deck(client):
    # the fixture has a V2 row; block policy turns it into an ERROR… but the
    # form doesn't expose the policy, so exercise the store contract directly
    from app.efficiency.analysis import ReportConfig, analyze
    import io
    a = analyze(io.BytesIO(build_eff_bytes()),
                ReportConfig(missing_row_policy="block"))
    token = eff_routes.EFF_REPORTS.put({"analysis": a, "pptx": None, "filename": "x"})
    page = client.get(f"/efficiency/{token}")
    assert page.status_code == 200
    assert "Deck not generated" in page.text
    assert client.get(f"/efficiency/{token}/deck.pptx").status_code == 404


def test_bad_file_is_422_not_500(client):
    r = client.post("/efficiency",
                    files={"report": ("junk.xlsx", b"not a zip", "application/zip")},
                    data={})
    assert r.status_code == 422


def test_malformed_ooxml_metadata_is_422_not_500(client):
    workbook = io.BytesIO()
    with zipfile.ZipFile(workbook, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types><Override")

    response = client.post(
        "/efficiency",
        files={"report": ("broken.xlsx", workbook.getvalue(),
                          "application/zip")},
        data={},
    )

    assert response.status_code == 422


def test_oversize_file_is_413(client, monkeypatch):
    monkeypatch.setattr(config, "EFF_MAX_UPLOAD_BYTES", 10)
    r = client.post("/efficiency",
                    files={"report": ("large.xlsx", b"x" * 11,
                                      "application/zip")},
                    data={})
    assert r.status_code == 413


def test_efficiency_limit_is_independent_of_reconciler_limit(client, monkeypatch):
    """The efficiency workbook has its own 40 MB budget — a file the
    reconciler's MAX_UPLOAD_BYTES would reject still analyzes here."""
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 10)   # reconciler-only cap
    r = _upload(client)
    assert r.status_code == 200


def test_bloat_in_other_sheets_does_not_reject_the_workbook(client, monkeypatch):
    """Real trackers carry extra tabs, styling, and pasted images that count
    toward the global cell/byte budgets; the streaming analyzer never reads
    them, so the efficiency gate must not reject on their account."""
    from openpyxl import load_workbook as _lw
    import io as _io
    monkeypatch.setattr(config, "MAX_XLSX_CELLS", 100)    # reconciler-strict
    wb = _lw(_io.BytesIO(build_eff_bytes()))
    junk = wb.create_sheet("YTD ARCHIVE")
    for row in range(1, 41):                              # 400 cells > cap
        for col in range(1, 11):
            junk.cell(row=row, column=col, value=f"x{row}-{col}")
    buf = _io.BytesIO()
    wb.save(buf)
    r = client.post(
        "/efficiency",
        files={"report": ("fat.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument"
                          ".spreadsheetml.sheet")},
        data={}, follow_redirects=True)
    assert r.status_code == 200
    assert "Basis: pooled" in r.text


def test_sheet_row_cap_rejects_with_a_clear_message(client, monkeypatch):
    monkeypatch.setattr(config, "EFF_MAX_SHEET_ROWS", 5)
    r = _upload(client)                                   # fixture has 44 rows
    assert r.status_code == 422
    assert "data rows" in r.text


def test_invalid_config_values_fall_back_to_defaults(client):
    r = _upload(client, basis="nonsense", tier_mode="nonsense", language="xx")
    assert r.status_code == 200
    assert "Basis: pooled" in r.text


def test_demo_image_shipped_and_shown_per_language(client):
    """The feature-demo slide images are synthetic-data renders committed to
    /static — both language variants must exist and be wired to the UI lang."""
    from pathlib import Path
    static = Path(main_mod.__file__).parent / "static"
    for lang in ("en", "zh"):
        assert (static / f"eff_demo_{lang}.jpg").stat().st_size > 10_000
    assert "/static/eff_demo_en.jpg" in client.get("/").text
    assert "/static/eff_demo_en.jpg" in client.get("/efficiency").text
    client.cookies.set("dmr_lang", "zh")
    assert "/static/eff_demo_zh.jpg" in client.get("/").text
    assert "/static/eff_demo_zh.jpg" in client.get("/efficiency").text


def test_store_cap_evicts_oldest():
    eff_routes.EFF_REPORTS.clear()
    tokens = [eff_routes.EFF_REPORTS.put({"analysis": {}, "pptx": None,
                                   "filename": str(i)})
              for i in range(eff_routes.EFF_REPORTS.max_entries + 3)]
    assert len(eff_routes.EFF_REPORTS) <= eff_routes.EFF_REPORTS.max_entries
    assert tokens[0] not in eff_routes.EFF_REPORTS      # oldest evicted
    assert tokens[-1] in eff_routes.EFF_REPORTS
    eff_routes.EFF_REPORTS.clear()


def test_chinese_filename_downloads_cleanly(client):
    """Report filenames here are usually Chinese; a raw filename= header is
    Latin-1 and 500'd. RFC 5987 filename* carries the real name."""
    r = client.post(
        "/efficiency",
        files={"report": ("小红书投放.xlsx", build_eff_bytes(),
                          "application/vnd.openxmlformats-officedocument"
                          ".spreadsheetml.sheet")},
        data={}, follow_redirects=True)
    assert r.status_code == 200
    token = next(iter(eff_routes.EFF_REPORTS))
    d = client.get(f"/efficiency/{token}/deck.pptx")
    assert d.status_code == 200
    assert d.content[:2] == b"PK"
    cd = d.headers["content-disposition"]
    assert "filename*=UTF-8''%E5%B0%8F%E7%BA%A2%E4%B9%A6" in cd
    assert 'filename="' in cd            # ASCII fallback for old clients
    assert cd.encode("latin-1")          # must be header-encodable
